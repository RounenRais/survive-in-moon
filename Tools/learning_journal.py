"""Helpers for writing Docs/Learning_Journal.docx. Usage: open the doc with python-docx, wrap it in Journal, add an entry, save."""
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ACCENT = RGBColor(0x1F, 0x4E, 0x79)

def shade(paragraph, fill):
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill)
    pPr.append(shd)

def border_left(paragraph, color):
    pPr = paragraph._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr"); left = OxmlElement("w:left")
    for k, v in (("val", "single"), ("sz", "24"), ("space", "8"), ("color", color)): left.set(qn("w:" + k), v)
    bdr.append(left); pPr.append(bdr)

def rich(paragraph, text):
    """**bold** and `code` inline markup."""
    import re
    for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
        if not part: continue
        if part.startswith("**"):
            r = paragraph.add_run(part[2:-2]); r.bold = True
        elif part.startswith("`"):
            r = paragraph.add_run(part[1:-1]); r.font.name = "Consolas"; r.font.size = Pt(10); r.font.color.rgb = RGBColor(0xA3, 0x1F, 0x34)
        else:
            paragraph.add_run(part)
    return paragraph

def new_numbering(doc, paragraph):
    """Creates a fresh numbering instance (restarting at 1) based on the List Number style's list."""
    numbering = doc.part.numbering_part.numbering_definitions._numbering
    style_numPr = doc.styles["List Number"].element.pPr.numPr
    abstract_id = None
    for num in numbering.findall(qn("w:num")):
        if num.get(qn("w:numId")) == style_numPr.numId.val.__str__():
            abstract_id = num.find(qn("w:abstractNumId")).get(qn("w:val"))
    ids = [int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))]
    new_id = max(ids) + 1
    num = OxmlElement("w:num"); num.set(qn("w:numId"), str(new_id))
    a = OxmlElement("w:abstractNumId"); a.set(qn("w:val"), abstract_id); num.append(a)
    o = OxmlElement("w:lvlOverride"); o.set(qn("w:ilvl"), "0")
    so = OxmlElement("w:startOverride"); so.set(qn("w:val"), "1"); o.append(so); num.append(o)
    numbering.append(num)
    return new_id

def set_num(paragraph, num_id):
    pPr = paragraph._p.get_or_add_pPr()
    numPr = OxmlElement("w:numPr")
    il = OxmlElement("w:ilvl"); il.set(qn("w:val"), "0"); numPr.append(il)
    ni = OxmlElement("w:numId"); ni.set(qn("w:val"), str(num_id)); numPr.append(ni)
    pPr.append(numPr)

class Journal:
    def __init__(self, doc): self.d = doc
    def h1(self, t): return self.d.add_heading(t, 1)
    def h2(self, t): return self.d.add_heading(t, 2)
    def h3(self, t): return self.d.add_heading(t, 3)
    def p(self, t): return rich(self.d.add_paragraph(), t)
    def bullet(self, t): return rich(self.d.add_paragraph(style="List Bullet"), t)
    def number(self, t):
        """Numbered item; a new list (starting from 1) begins automatically after any other paragraph."""
        last = self.d.paragraphs[-1] if self.d.paragraphs else None
        para = self.d.add_paragraph(style="List Number")
        if last is None or last.style.name != "List Number" or not hasattr(self, "_num_id"):
            self._num_id = new_numbering(self.d, para)
        set_num(para, self._num_id)
        return rich(para, t)
    def box(self, title, text, fill="EAF2FB", color="2E75B6"):
        para = self.d.add_paragraph(); shade(para, fill); border_left(para, color)
        r = para.add_run(title + "  "); r.bold = True; r.font.color.rgb = RGBColor.from_string(color)
        rich(para, text); para.paragraph_format.space_after = Pt(8)
        return para
    def concept(self, name, text): return self.box("📘 Yeni kavram: " + name, text)
    def tip(self, text): return self.box("💡 İpucu:", text, "FFF6DD", "C89B00")
    def warn(self, text): return self.box("⚠️ Dikkat:", text, "FDECEC", "C0392B")
    def code(self, text):
        for line in text.strip("\n").split("\n"):
            para = self.d.add_paragraph(); shade(para, "F3F3F3")
            para.paragraph_format.space_after = Pt(0); para.paragraph_format.left_indent = Cm(0.4)
            r = para.add_run(line if line else " "); r.font.name = "Consolas"; r.font.size = Pt(9.5)
        self.d.add_paragraph()
    def table(self, header, rows, widths=None):
        t = self.d.add_table(rows=1, cols=len(header)); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, h in enumerate(header):
            c = t.rows[0].cells[i]; c.text = ""; r = c.paragraphs[0].add_run(h); r.bold = True; r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            c.paragraphs[0].paragraph_format.keep_with_next = True
            tcPr = c._tc.get_or_add_tcPr(); shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:fill"), "1F4E79"); tcPr.append(shd)
        trPr = t.rows[0]._tr.get_or_add_trPr(); hdr = OxmlElement("w:tblHeader"); hdr.set(qn("w:val"), "true"); trPr.append(hdr)
        for row in rows:
            cells = t.add_row().cells
            for i, v in enumerate(row): cells[i].text = ""; rich(cells[i].paragraphs[0], v)
        if widths:
            for row in t.rows:
                for i, w in enumerate(widths): row.cells[i].width = Cm(w)
        self.d.add_paragraph()
        return t
    def image(self, path, width_cm=16, caption=None):
        self.d.add_picture(path, width=Cm(width_cm)); self.d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        if caption:
            c = self.d.add_paragraph(); c.alignment = WD_ALIGN_PARAGRAPH.CENTER; r = c.add_run(caption); r.italic = True; r.font.size = Pt(9)
    def entry(self, number, date, title):
        self.d.add_page_break()
        h = self.h1(f"Kayıt {number} — {title}")
        m = self.d.add_paragraph(); r = m.add_run(f"Tarih: {date}"); r.italic = True; r.font.color.rgb = RGBColor(0x70, 0x70, 0x70)
        return h

def setup_styles(doc):
    st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(11)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    for name, size in (("Heading 1", 18), ("Heading 2", 14), ("Heading 3", 12)):
        s = doc.styles[name]; s.font.name = "Calibri"; s.font.size = Pt(size); s.font.color.rgb = ACCENT
    for sec in doc.sections:
        sec.left_margin = sec.right_margin = Cm(2.2); sec.top_margin = sec.bottom_margin = Cm(2)
